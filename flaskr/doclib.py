from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.shared import Inches, Pt


def generate(resume_json):
    """Returns a doc of the given resume JSON file."""
    # Prepare output file stream
    file_stream = BytesIO()

    # Create document and save it to files stream
    document = Document()
    set_layout(document)
    set_defaults(document)
    add_header(document, resume_json["header"])
    add_sections(document, resume_json["sections"])

    # Save document to files stream
    document.save(file_stream)
    file_stream.seek(0)

    return file_stream


def set_defaults(document):
    """Sets a default font and size for the document."""
    style = document.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.paragraph_format.space_after = Pt(0)


def set_layout(document):
    """Sets the page layout of a given Document object."""
    for section in document.sections:
        # Narrow margins
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)

        # A4 size (8.27 x 11.69 inches)
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)


def add_header(document, header):
    """Adds the header to the given resume document."""
    # Get header parameters
    name = header["name"]
    email = header["email"]
    mobile = header["mobile"]
    portfolio = header["portfolio"]
    linkedin = header["linkedin"]

    # Name paragraph
    paragraph = document.add_paragraph()
    run = paragraph.add_run(name)
    run.bold = True
    run.font.size = Pt(16)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Email and mobile paragraph
    paragraph = document.add_paragraph(f"Email: {email}, Mobile: {mobile}")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Portfolio and LinkedIn paragraph
    paragraph = document.add_paragraph(f"Portfolio: {portfolio}, LinkedIn: {linkedin}")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def add_sections(document, sections):
    """Adds sections to the given resume document."""
    for index, section in enumerate(sections):
        add_section(document, section)

        # Add a newline if the section is not the last
        if index != len(sections) - 1:
            add_newline(document)


def add_section(document, section):
    """Add a particular section to the given resume document."""
    # Add section title
    title = section["title"].upper()
    paragraph = document.add_paragraph()
    run = paragraph.add_run(title)
    run.bold = True
    run.font.size = Pt(12)

    # Add section contents
    if "entries" in section:
        entries = section["entries"]
        add_entries(document, entries)
    elif "cslist" in section:
        cslist = section["cslist"]
        add_cslist(document, cslist)
    elif "certifications" in section:
        certifications = section["certifications"]
        add_certifications(document, certifications)


def add_entries(document, entries):
    """Add entries (of a section) to the given resume document."""
    for index, entry in enumerate(entries):
        add_entry(document, entry)

        # Add a newline if entry is not the last
        if index != len(entries) - 1:
            add_newline(document)


def add_entry(document, entry):
    """Add an entry of a section to the given resume document."""
    # Get entry parameters
    name = entry["name"]
    dates = entry["date(s)"]

    # Name and dates paragraph
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.tab_stops.add_tab_stop(
        Inches(7.27),  # right margin position
        alignment=WD_TAB_ALIGNMENT.RIGHT,
    )
    run = paragraph.add_run(f"{name}\t{dates}")
    run.bold = True

    # Company paragraph (if present)
    if "company" in entry:
        company = entry["company"]
        paragraph = document.add_paragraph()
        run = paragraph.add_run(company)
        run.bold = True

    # Bulletpoints
    bulletlist = entry["bulletlist"]
    for bulletpoint in bulletlist:
        document.add_paragraph(bulletpoint, style="List Bullet")


def add_cslist(document, cslist):
    """Add a comma-separated list of a section to the given resume document."""
    document.add_paragraph(", ".join(cslist))


def add_certifications(document, certifications):
    """Add certifications (of a section) to the given resume document."""
    for certification in certifications:
        add_certification(document, certification)


def add_certification(document, certification):
    """Add a given certification of a section to the given resume document."""
    # Get paragraph parameters
    name = certification["name"]
    institution = certification["institution"]
    dates = certification["date(s)"]

    # Create paragraph
    paragraph = document.add_paragraph("", style="List Bullet")
    run = paragraph.add_run(name)
    run.bold = True
    paragraph.add_run(f", {institution} - {dates}")


def add_newline(document):
    """Add a newline to the given resume document."""
    document.add_paragraph()
